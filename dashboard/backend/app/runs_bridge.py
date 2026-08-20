from __future__ import annotations

import base64
from pathlib import Path

from fastapi.responses import HTMLResponse
import json as _json_mod

from fastapi import Depends, APIRouter, Request
from fastapi.responses import PlainTextResponse
from fastapi.responses import FileResponse, JSONResponse, Response

from . import calibration, dataset, handover, mobility, provenance, ran_bridge, report, runs_store, sionna_bridge, stability

from .auth import require_user  # noqa: E402

router = APIRouter(prefix="/api/recordings", tags=["recordings"],
                   dependencies=[Depends(require_user)])


def _bad(msg: str, code: int = 400):
    return JSONResponse({"ok": False, "error": msg}, status_code=code)


async def _json(req: Request) -> dict:
    try:
        return await req.json()
    except Exception:
        return {}


def _sources() -> dict:
    return {
        "kpi": lambda: ran_bridge.kpi(samples=300),
        "live": lambda: sionna_bridge.LIVE.status(),
        "ue_sinr": lambda: calibration.ue_sinr_sample(ran_bridge.UE_LOG,
                                                      require_fresh=True),
        "logs": _logs_now,
    }


def _logs_now() -> dict:
    out: dict = {}
    for name in ran_bridge.LOG_SOURCES:
        r = ran_bridge.logs(source=name, lines=4000)
        if isinstance(r, dict) and r.get("lines"):
            out[name] = r["lines"]
    st = sionna_bridge.LIVE.status()
    if st.get("log"):
        out["sionna"] = st["log"]
    return out


@router.get("")
@router.get("/")
def runs_list(limit: int = 50, kind: str = ""):
    k = kind if kind == "sim" else None
    return {"ok": True, "recording": runs_store.RECORDER.run_id,
            "runs": runs_store.list_runs(limit, k),
            "counts": runs_store.count_runs()}


@router.get("/open")
def run_open():
    rid = runs_store.open_run_id()
    return {"ok": True, "id": rid,
            "live": bool(rid and runs_store.RECORDER.run_id == rid
                         and runs_store.RECORDER.is_recording())}


@router.post("/start")
async def run_start(req: Request):
    body = await _json(req)
    if runs_store.RECORDER.is_recording():
        return _bad(f"a run is already recording ({runs_store.RECORDER.run_id}) — stop it first", 409)
    runs_store.close_orphans()
    scene_id = body.get("scene_id")
    scene_name = None
    if scene_id:
        d = sionna_bridge._scene_dir(str(scene_id))
        if d:
            scene_name = (sionna_bridge._read_meta(d) or {}).get("name")
    d = sionna_bridge._scene_dir(str(scene_id)) if scene_id else None
    params = {
        "provenance": provenance.run_provenance(d),
        "du_config": body.get("du_config"),
        "ue": sionna_bridge.ue_params(sionna_bridge._scene_dir(str(scene_id))) if scene_id
              and sionna_bridge._scene_dir(str(scene_id)) else None,
        "sionna": sionna_bridge.sionna_params(sionna_bridge._scene_dir(str(scene_id))) if scene_id
                  and sionna_bridge._scene_dir(str(scene_id)) else None,
    }
    return runs_store.start_run(scene_id, scene_name, body.get("label"),
                                body.get("du_config"), params, _sources())


@router.post("/{run_id}/frame")
async def run_frame(run_id: str, req: Request):
    body = await _json(req)
    src = str(body.get("image") or "")
    if "," in src:
        src = src.split(",", 1)[1]
    try:
        data = base64.b64decode(src, validate=True)
    except Exception:
        return _bad("image must be a base64 data URL")
    if not data:
        return _bad("empty image")
    try:
        n = runs_store.add_frame(run_id, data)
    except FileNotFoundError:
        return _bad("unknown run_id", 404)
    return {"ok": True, "frames": n}


@router.post("/{run_id}/stop")
async def run_stop(run_id: str, req: Request):
    body = await _json(req)
    try:
        return runs_store.stop_run(run_id, _logs_now(), body.get("extra"))
    except FileNotFoundError:
        return _bad("unknown run_id", 404)


@router.get("/policies")
def list_policies():
    return {"ok": True, "policies": sorted(mobility.POLICIES),
            "defaults": {"floor_db": mobility.DEFAULT_FLOOR_DB,
                         "ceiling_db": mobility.DEFAULT_CEILING_DB},
            "note": ("patrol is the open-loop baseline: it ignores every measurement. A "
                     "closed-loop policy that produces the patrol track reacted to nothing.")}


@router.get("/{run_id}")
def run_detail(run_id: str):
    r = runs_store.get_run(run_id)
    if not r:
        return _bad("unknown run_id", 404)
    return {"ok": True, "run": r}


@router.get("/{run_id}/log/{source}")
def run_log(run_id: str, source: str):
    lines = runs_store.read_log(run_id, source)
    if lines is None:
        return _bad("no such log for this run", 404)
    return {"ok": True, "source": source, "lines": lines}


@router.get("/{run_id}/video")
def run_video(run_id: str):
    p = runs_store.run_dir(run_id) / "movement.mp4"
    if not p.exists():
        return _bad("this run has no video", 404)
    return FileResponse(p, media_type="video/mp4")


@router.get("/{run_id}/frame/{n}")
def run_frame_image(run_id: str, n: int):
    p = runs_store.run_dir(run_id) / "frames" / f"{int(n):05d}.jpg"
    if not p.exists():
        return _bad("no such frame", 404)
    return FileResponse(p, media_type="image/jpeg")


@router.delete("/{run_id}")
def run_delete(run_id: str):
    if runs_store.RECORDER.run_id == run_id:
        return _bad("this run is still recording — stop it first", 409)
    return {"ok": True, "deleted": runs_store.delete_run(run_id)}


@router.get("/{run_id}/calibration")
def run_calibration(run_id: str, tolerance_s: float = calibration.DEFAULT_TOLERANCE_S,
                    plot: bool = True):
    run = runs_store.get_run(run_id)
    if not run:
        return _bad("unknown run_id", 404)
    d = runs_store.run_dir(run_id)
    if not d.is_dir():
        return _bad("this run has no recorded directory", 404)

    report = calibration.calibrate(d, run, tolerance_s=float(tolerance_s))
    if plot:
        for c in report["comparisons"]:
            c["svg"] = calibration.svg_scatter(
                c["pairs"], c["pairs_fit"],
                x_label="traced path gain (dB)",
                y_label=f"{c['label']} ({c['unit']})")
    return {"ok": True, "run_id": run_id, "scene_id": run.get("scene_id"),
            "scene_name": run.get("scene_name"), **report}


@router.get("/{run_id}/dataset")
def run_dataset(run_id: str, fmt: str = "json",
                tolerance_s: float = calibration.DEFAULT_TOLERANCE_S,
                lag_s: float = None):
    run = runs_store.get_run(run_id)
    if not run:
        return _bad("unknown run_id", 404)
    d = runs_store.run_dir(run_id)
    if not d.is_dir():
        return _bad("this run has no recorded directory", 404)

    scene_dir = None
    if run.get("scene_id"):
        scene_dir = sionna_bridge._scene_dir(str(run["scene_id"]))

    built = dataset.build_rows(d, run, tolerance_s=float(tolerance_s),
                               lag_s=lag_s, scene_dir=scene_dir,
                               enriched=_enriched(run_id))
    rows = built["rows"]

    fmt = (fmt or "json").lower()
    if fmt == "csv":
        return PlainTextResponse(dataset.to_csv(rows), media_type="text/csv")
    if fmt == "xlsx":
        """The same rows, laid out to be read rather than parsed.

        Excel truncates a header it has not been told to widen, and 34 columns of telemetry are
        unreadable without a frozen top row. None of that belongs in the CSV, which tools read,
        so it lives in its own format instead.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter
        except Exception:
            return _bad("xlsx export needs openpyxl on the server; fmt=csv still works", 501)
        import datetime as _dt
        import io as _io

        wb = Workbook()
        ws = wb.active
        ws.title = "KPIs"
        cols = list(dataset.ALL_KEYS)
        ws.append(cols)
        for r in rows:
            ws.append([r.get(k) for k in cols])

        if "t" in cols:
            ti = cols.index("t") + 1
            for row in ws.iter_rows(min_row=2, min_col=ti, max_col=ti):
                for c in row:
                    if isinstance(c.value, (int, float)):
                        c.value = _dt.datetime.fromtimestamp(c.value)
                        c.number_format = "yyyy-mm-dd hh:mm:ss.000"

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, name in enumerate(cols, start=1):
            longest = len(str(name))
            for c in ws[get_column_letter(i)][1:400]:
                if c.value is not None:
                    longest = max(longest, len(str(c.value)))
            ws.column_dimensions[get_column_letter(i)].width = min(34, max(10, longest + 2))
            ws.cell(row=1, column=i).font = Font(bold=True)

        buf = _io.BytesIO()
        wb.save(buf)
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="%s_kpis.xlsx"' % run_id},
        )
    if fmt == "jsonl":
        return PlainTextResponse(dataset.to_jsonl(rows),
                                 media_type="application/x-ndjson")
    if fmt != "json":
        return _bad("fmt must be json, csv, jsonl or xlsx")

    return {
        "ok": True, "run_id": run_id, "scene_id": run.get("scene_id"),
        "scene_name": run.get("scene_name"),
        "rows": len(rows),
        "dropped": built["dropped"],
        "lag_s": built["lag_s"], "lag_source": built["lag_source"],
        "lag_r2": built["lag_r2"], "tolerance_s": built["tolerance_s"],
        "ue_tolerance_s": built["ue_tolerance_s"],
        "enriched": built["enriched"],
        "enriched_positions": built["enriched_positions"],
        "uplink": built["uplink"],
        "schema": dataset.schema(),
        "split": dataset.split_advice(rows),
        "coverage": dataset.coverage(rows),
        "data": rows,
    }


ENRICHED_NAME = "enriched.json"


def _enriched(run_id: str):
    import json as _json
    p = runs_store.run_dir(run_id) / ENRICHED_NAME
    try:
        return _json.loads(p.read_text())
    except (OSError, ValueError):
        return None


@router.get("/{run_id}/enrichment")
def run_enrichment(run_id: str):
    run = runs_store.get_run(run_id)
    if not run:
        return _bad("unknown run_id", 404)
    if not runs_store.run_dir(run_id).is_dir():
        return _bad("this run has no recorded directory", 404)

    en = _enriched(run_id)
    if not en:
        return {"ok": True, "run_id": run_id, "enriched": False,
                "how": f"POST /api/recordings/{run_id}/enrich"}

    rows = list((en.get("positions") or {}).values())
    if not rows:
        return {"ok": True, "run_id": run_id, "enriched": True, "positions_traced": 0,
                "note": "the enrichment recorded no positions"}

    def span(key):
        vals = [r.get(key) for r in rows if r.get(key) is not None]
        if not vals:
            return None
        return {"min": min(vals), "max": max(vals), "n": len(vals)}

    los = [bool(r.get("los")) for r in rows if r.get("los") is not None]
    counts = {}
    for kind in ("specular", "refraction", "diffraction", "diffuse"):
        key = "n_" + kind
        vals = [r.get(key) for r in rows if r.get(key) is not None]
        if vals:
            counts[kind] = {"min": min(vals), "max": max(vals), "total": sum(vals)}

    cells = sorted({name for r in rows for name in (r.get("per_tx") or {})})

    return {
        "ok": True,
        "run_id": run_id,
        "enriched": True,
        "positions_traced": en.get("distinct"),
        "updates_covered": en.get("requested"),
        "max_depth": en.get("max_depth"),
        "rx_height_m": en.get("rx_height_m"),
        "cells": cells,
        "los": {"true": sum(1 for v in los if v), "of": len(los)} if los else None,
        "dist_m": span("dist_m"),
        "n_paths_traced": span("n_paths_traced"),
        "rms_delay_spread_s": span("rms_delay_spread_s"),
        "first_arrival_s": span("first_arrival_s"),
        "strongest_share": span("strongest_share"),
        "strongest_bounces": span("strongest_bounces"),
        "interactions": counts,
        "positions": en.get("positions"),
    }


@router.post("/{run_id}/enrich")
def run_enrich(run_id: str):
    import json as _json
    run = runs_store.get_run(run_id)
    if not run:
        return _bad("unknown run_id", 404)
    d = runs_store.run_dir(run_id)
    if not d.is_dir():
        return _bad("this run has no recorded directory", 404)

    scene_dir = sionna_bridge._scene_dir(str(run.get("scene_id") or ""))
    if not scene_dir:
        return _bad("this run has no scene to re-trace against", 400)
    xml = scene_dir / "scene" / "scene.xml"
    pkg = scene_dir / "package.json"
    if not xml.exists() or not pkg.exists():
        return _bad("the scene this run used is no longer built", 400)

    positions = [s["pos"] for s in calibration.load_series(d / "telemetry.jsonl")
                 if s.get("pos")]
    if not positions:
        return _bad("this run recorded no positions", 400)

    pos_file = d / "_positions.json"
    pos_file.write_text(_json.dumps(positions))
    out_file = d / ENRICHED_NAME

    cmd = [sionna_bridge.SIONNA_PY,
           str(sionna_bridge.SCENE_BUILDER / "enrich.py"), str(pkg),
           "--scene", str(xml), "--positions", str(pos_file), "--out", str(out_file),
           "--json"]
    cmd += sionna_bridge._solver_flags(sionna_bridge.sionna_params(scene_dir))

    def work(job):
        try:
            tail = sionna_bridge._stream(job, cmd, 1800)
        finally:
            pos_file.unlink(missing_ok=True)
        try:
            summary = _json.loads(tail.strip().splitlines()[-1])
        except (ValueError, IndexError):
            raise RuntimeError("the enricher produced no result; see the job log")
        return {"run_id": run_id, **summary}

    job = sionna_bridge.submit("run-enrich", f"enrich · {run_id}", work,
                               scene_id=run.get("scene_id"))
    return sionna_bridge._job_response(job, {"wait": True}, extra={"run_id": run_id})


@router.get("/{run_id}/policy")
def run_policy(run_id: str, policy: str = "seek_coverage", floor_db: float = None,
               ceiling_db: float = None, seed: int = 0, start_index: int = 0):
    run = runs_store.get_run(run_id)
    if not run:
        return _bad("unknown run_id", 404)
    if policy not in mobility.POLICIES:
        return _bad(f"policy must be one of {sorted(mobility.POLICIES)}")
    d = runs_store.run_dir(run_id)

    ue = calibration.load_series(d / "ue_kpi.jsonl")
    tel = calibration.load_series(d / "telemetry.jsonl")
    if not tel:
        return _bad("this run recorded no channel updates", 400)

    paired = calibration.align_nearest(tel, ue, left_key="seq", right_key="sinr",
                                       tolerance_s=dataset.UE_TOLERANCE_S)
    readings = [{"ue_sinr_db": p["y"]} for p in paired]
    n_positions = len({tuple(s["pos"]) for s in tel if s.get("pos")}) or len(tel)

    params = {}
    if floor_db is not None:
        params["floor_db"] = float(floor_db)
    if ceiling_db is not None:
        params["ceiling_db"] = float(ceiling_db)

    out = mobility.run_policy(policy, readings, n_positions=n_positions,
                              start_index=int(start_index), params=params, seed=int(seed))
    out["compared_to_patrol"] = mobility.differs_from_patrol(out, n_positions)
    out["readings"] = len(readings)
    out["n_positions"] = n_positions
    out["run_id"] = run_id
    out["ok"] = True
    return out


@router.get("/{run_id}/report")
def run_report(run_id: str, fmt: str = "html",
               tolerance_s: float = calibration.DEFAULT_TOLERANCE_S):
    try:
        rep = report.build(run_id, tolerance_s=float(tolerance_s))
    except FileNotFoundError:
        return _bad("unknown run_id", 404)
    if fmt == "json":
        return {"ok": True, **rep}
    return HTMLResponse(report.render_html(rep))


@router.get("/{run_id}/stability")
def run_stability(run_id: str, live: bool = True):
    run = runs_store.get_run(run_id)
    if not run:
        return _bad("unknown run_id", 404)
    d = runs_store.run_dir(run_id)
    if not d.is_dir():
        return _bad("this run has no recorded directory", 404)
    return {"ok": True, "run_id": run_id, **stability.diagnose(d, live=bool(live))}


@router.get("/{run_id}/handover")
def run_handover(run_id: str, hysteresis_db: float = handover.DEFAULT_HYSTERESIS_DB,
                 offset_db: float = handover.DEFAULT_OFFSET_DB,
                 time_to_trigger_s: float = handover.DEFAULT_TIME_TO_TRIGGER_S):
    run = runs_store.get_run(run_id)
    if not run:
        return _bad("unknown run_id", 404)
    d = runs_store.run_dir(run_id)
    if not d.is_dir():
        return _bad("this run has no recorded directory", 404)

    enriched_path = d / ENRICHED_NAME
    if not enriched_path.exists():
        return _bad(f"this run is not enriched yet — POST /api/recordings/{run_id}/enrich", 409)
    try:
        enriched = _json_mod.loads(enriched_path.read_text())
    except (OSError, ValueError) as exc:
        return _bad(f"could not read the enrichment: {exc}", 500)

    positions = enriched.get("positions") or {}
    cells = sorted({name for row in positions.values() for name in (row.get("per_tx") or {})})
    if len(cells) < 2:
        return {"ok": True, "run_id": run_id, "cells": cells,
                "predicted": {"events": [], "n": 0},
                "note": ("this scene has %d transmitter(s), so there is no neighbour to hand "
                         "over to and no A3 event can exist. Add a second gNB to the scene to "
                         "ask this question." % len(cells))}

    series = []
    for sample in calibration.load_series(d / "telemetry.jsonl"):
        pos, t = sample.get("pos"), calibration.to_epoch(sample.get("t"))
        if not pos or t is None:
            continue
        row = positions.get(f"{float(pos[0]):.2f},{float(pos[1]):.2f}")
        if not row:
            continue
        rss = {name: v.get("rss_dbm") for name, v in (row.get("per_tx") or {}).items()
               if v.get("rss_dbm") is not None}
        if rss:
            series.append({"t": t, "rss": rss})

    if len(series) < 2:
        return _bad("not enough traced positions in this run to run the A3 rule", 409)

    predicted = handover.a3_events(series, hysteresis_db=float(hysteresis_db),
                                   offset_db=float(offset_db),
                                   time_to_trigger_s=float(time_to_trigger_s))
    actual = handover.parse_actual(runs_store.read_log(run_id, "cu") or [])
    out = {
        "ok": True, "run_id": run_id, "cells": cells, "samples": len(series),
        "params": {"hysteresis_db": hysteresis_db, "offset_db": offset_db,
                   "time_to_trigger_s": time_to_trigger_s},
        "predicted": predicted,
        "actual": actual,
        "naive_switches": len(handover.naive_switches(series)),
        "ping_pongs": handover.ping_pongs(predicted.get("events") or []),
    }
    if actual.get("events"):
        out["comparison"] = handover.compare(predicted.get("events") or [], actual["events"])
    return out
